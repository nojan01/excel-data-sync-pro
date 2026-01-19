#!/usr/bin/env python3
"""Prüfe Tables in Original-Datei"""
from openpyxl import load_workbook

INPUT_FILE = '/Users/nojan/Desktop/2025-08-31 DEFENCE&SPACE MVMS Master Asset List GER.xlsx'
SHEET_NAME = 'DEFENCE&SPACE Aug-2025'

print(f"Lade: {INPUT_FILE}")
wb = load_workbook(INPUT_FILE)
ws = wb[SHEET_NAME]

print(f"\n=== Tables im Sheet ===")
for table in ws.tables.values():
    print(f"Table: {table.name}")
    print(f"  Range: {table.ref}")
    print(f"  Display Name: {table.displayName}")
    print(f"  Table Style: {table.tableStyleInfo}")
    if table.tableStyleInfo:
        print(f"    - name: {table.tableStyleInfo.name}")
        print(f"    - showRowStripes: {table.tableStyleInfo.showRowStripes}")
        print(f"    - showColumnStripes: {table.tableStyleInfo.showColumnStripes}")
        print(f"    - showFirstColumn: {table.tableStyleInfo.showFirstColumn}")
        print(f"    - showLastColumn: {table.tableStyleInfo.showLastColumn}")
    print(f"  AutoFilter: {table.autoFilter}")
    print(f"  Columns: {len(table.tableColumns)} Spalten")
    for i, col in enumerate(table.tableColumns[:5]):
        print(f"    [{i}] {col.name}")
    if len(table.tableColumns) > 5:
        print(f"    ... und {len(table.tableColumns) - 5} weitere")

wb.close()
