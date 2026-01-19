#!/usr/bin/env python3
"""Prüfe Table Columns in der gespeicherten Datei"""

from openpyxl import load_workbook

try:
    wb = load_workbook('/Users/nojan/Desktop/TEST_WRITE_SHEET.xlsx')
    ws = wb['DEFENCE&SPACE Aug-2025']
    
    print('=== TABLE COLUMNS CHECK ===')
    
    for table in ws.tables.values():
        print(f'\nTable: {table.name}')
        print(f'Ref: {table.ref}')
        print(f'Columns count: {len(table.tableColumns)}')
        
        print('\nErste 15 Columns:')
        for i, col in enumerate(table.tableColumns[:15]):
            print(f'  {i+1}: id={col.id}, name="{col.name}"')
        
        # Prüfe ob alle Columns einen Namen haben
        print('\nPrüfe auf None-Namen:')
        for i, col in enumerate(table.tableColumns):
            if col.name is None:
                print(f'  ❌ Column {i+1} (id={col.id}) hat name=None!')
        
    wb.close()
except Exception as e:
    print(f'Fehler: {e}')
